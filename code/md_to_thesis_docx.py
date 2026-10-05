import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


WORKSPACE = Path(__file__).resolve().parents[1]
SRC = WORKSPACE / "논문_통합본_수정본.md"
OUT_DIR = WORKSPACE / "outputs" / "hwp_work"
OUT = OUT_DIR / "논문_통합본_한글작업본.docx"


TITLE = "배전설비 자산투자계획(AIP)을 위한\nPriority Number 기반 이종설비\n5개년 포트폴리오 최적화 모델"


def set_run_font(run, font_name: str, size_pt: float | None = None, bold: bool | None = None):
    """한글 글꼴이 렌더러에서 빠지지 않도록 rFonts를 함께 지정한다."""
    run.font.name = font_name
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        r_fonts.set(qn(key), font_name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def set_paragraph_format(paragraph, *, align=None, before=0, after=0, line=2.0, first_line=False):
    """논문양식의 본문 줄간격과 문단 간격을 적용한다."""
    fmt = paragraph.paragraph_format
    if align is not None:
        paragraph.alignment = align
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Cm(0.7)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_page_number(section):
    """꼬리말 중앙에 PAGE 필드를 넣는다."""
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, line=1.0)
    run = p.add_run()
    set_run_font(run, "Batang", 9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_section(section):
    """논문양식.pdf의 A4 여백 기준을 적용한다."""
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(35)
    section.left_margin = Mm(35)
    section.right_margin = Mm(30)
    section.bottom_margin = Mm(25)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(15)


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Batang"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Batang")
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 2.0
    normal.paragraph_format.space_after = Pt(0)

    for name, size, before, after, align in [
        ("Heading 1", 16, 18, 12, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 13, 12, 6, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 12, 10, 4, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        st = styles[name]
        st.font.name = "Batang"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Batang")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 2.0
        st.paragraph_format.alignment = align


def clean_inline(text: str) -> str:
    text = text.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    text = re.sub(r"<!--.*?-->", "", text)
    text = text.replace("**", "")
    text = text.replace("`", "")
    return text.strip()


def add_plain_paragraph(doc: Document, text: str, *, align=None, first_line=True, font="Batang", size=11, bold=False):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=align, before=0, after=0, line=2.0, first_line=first_line)
    for idx, part in enumerate(text.split("\n")):
        if idx:
            p.add_run().add_break()
        run = p.add_run(part)
        set_run_font(run, font, size, bold)
    return p


def add_heading(doc: Document, level: int, text: str):
    style = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(clean_inline(text))
    set_run_font(run, "Batang", 16 if level == 1 else 13 if level == 2 else 12, True)
    return p


def add_cover_page(doc: Document, cover_type: str):
    """표지와 내표지를 논문양식에 맞게 별도 구성한다."""
    if cover_type == "cover":
        spacer_lines = [("학위논문", 16, False), ("", 12, False)]
        for text, size, bold in spacer_lines:
            add_plain_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, size=size, bold=bold)
        p = doc.add_paragraph()
        set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=36, after=36, line=1.6)
        run = p.add_run(TITLE)
        set_run_font(run, "HY견명조", 18, True)
        for text in ["박사학위논문", "", "국립목포대학교 대학원", "전기공학과", "[성 명]", "", "20   년    월"]:
            add_plain_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, size=14 if text else 11)
        doc.add_page_break()
        return

    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=36, after=36, line=1.6)
    run = p.add_run(TITLE)
    set_run_font(run, "HY견명조", 18, True)
    for text in [
        "이 논문을 공학 박사학위 논문으로 제출함",
        "",
        "국립목포대학교 대학원 전기공학과",
        "[성 명] (지도교수 : [지도교수명])",
        "",
        "[성 명]의 공학 박사학위 논문을 인준함",
        "",
        "심사위원장 　　　　　　 (인)",
        "심사위원 　　　　　　　 (인)",
        "심사위원 　　　　　　　 (인)",
        "심사위원 　　　　　　　 (인)",
        "심사위원 　　　　　　　 (인)",
        "",
        "20   년    월",
    ]:
        add_plain_paragraph(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, size=12 if text else 11)
    doc.add_page_break()


def collect_toc(lines: list[str]) -> list[tuple[int, str]]:
    items = []
    for line in lines:
        m = re.match(r"^(#{1,3})\s+(.+)$", line.strip())
        if not m:
            continue
        text = clean_inline(m.group(2))
        if text in {"목 차"} or text.startswith("배전설비 자산투자계획"):
            continue
        items.append((len(m.group(1)), text))
    return items


def add_static_toc(doc: Document, headings: list[tuple[int, str]]):
    add_heading(doc, 1, "목 차")
    for level, text in headings:
        p = doc.add_paragraph()
        set_paragraph_format(p, before=0, after=0, line=1.6, first_line=False)
        p.paragraph_format.left_indent = Cm(0.7 * (level - 1))
        run = p.add_run(text)
        set_run_font(run, "Batang", 11 if level == 1 else 10.5, False)
    doc.add_page_break()


def parse_table(lines: list[str], start: int):
    table_lines = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        table_lines.append(lines[idx].strip())
        idx += 1
    rows = []
    for line in table_lines:
        cells = [clean_inline(c.strip()) for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    return rows, idx


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            text = row[c_idx] if c_idx < len(row) else ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 or len(text) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_format(p, before=0, after=0, line=1.25, first_line=False)
            run = p.add_run(text)
            set_run_font(run, "Batang", 9, r_idx == 0)
            if r_idx == 0:
                set_cell_shading(cell, "F2F2F2")
    doc.add_paragraph()


def add_list_item(doc: Document, raw: str):
    indent = len(raw) - len(raw.lstrip(" "))
    stripped = raw.strip()
    bullet = re.sub(r"^[-*]\s+", "", stripped)
    numbered = re.sub(r"^\d+\.\s+", "", stripped)
    text = numbered if numbered != stripped else bullet
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=0, line=2.0, first_line=False)
    p.paragraph_format.left_indent = Cm(0.7 + indent * 0.08)
    p.paragraph_format.first_line_indent = Cm(-0.3)
    prefix = "• " if numbered == stripped else f"{stripped.split('.')[0]}. "
    run = p.add_run(prefix + clean_inline(text))
    set_run_font(run, "Batang", 11, False)


def add_code_block(doc: Document, code: str):
    for line in code.splitlines():
        p = doc.add_paragraph()
        set_paragraph_format(p, before=0, after=0, line=1.2, first_line=False)
        p.paragraph_format.left_indent = Cm(0.8)
        run = p.add_run(line)
        set_run_font(run, "Courier New", 9, False)


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    lines = SRC.read_text(encoding="utf-8-sig").splitlines()

    doc = Document()
    configure_section(doc.sections[0])
    configure_styles(doc)

    add_cover_page(doc, "cover")
    add_cover_page(doc, "inner")
    add_static_toc(doc, collect_toc(lines))

    in_code = False
    code_buf = []
    skip_front_matter = True
    idx = 0
    while idx < len(lines):
        raw = lines[idx]
        line = raw.rstrip()
        stripped = line.strip()

        # 원본 Markdown의 표지·내표지·목차 블록은 이미 별도 구성했으므로 국문초록부터 본문 변환을 시작한다.
        if skip_front_matter:
            if stripped == "# 국문초록":
                skip_front_matter = False
            else:
                idx += 1
                continue

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue

        if in_code:
            code_buf.append(line)
            idx += 1
            continue

        if not stripped:
            idx += 1
            continue

        if stripped == "---":
            idx += 1
            continue

        if stripped.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_table(doc, rows)
            continue

        if stripped.startswith(">"):
            # 작업본 안내문은 최종 문서에 넣지 않는다.
            idx += 1
            continue

        if stripped.startswith("<") and stripped.endswith(">"):
            idx += 1
            continue

        m = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if m:
            level = min(len(m.group(1)), 3)
            text = clean_inline(m.group(2))
            if level == 1 and text.startswith("제") and not text.startswith("제1장"):
                doc.add_page_break()
            elif level == 1 and text in {"참고문헌", "Abstract", "부록"}:
                doc.add_page_break()
            add_heading(doc, level, text)
            idx += 1
            continue

        if re.match(r"^\s*[-*]\s+", line) or re.match(r"^\s*\d+\.\s+", line):
            add_list_item(doc, line)
            idx += 1
            continue

        add_plain_paragraph(doc, clean_inline(line), first_line=True)
        idx += 1

    doc.save(OUT)
    print(str(OUT))


if __name__ == "__main__":
    build()
