# -*- coding: utf-8 -*-
"""
최신 논문_통합본_수정본.md를 한글 조판 전 단계용 DOCX로 변환한다.

설계 원칙
- 원본 Markdown은 변경하지 않는다.
- 실제 존재하는 그림 파일은 본문 위치에 삽입한다.
- 아직 산출되지 않은 그림/민감도 결과는 본문 텍스트 표기를 유지한다.
- HWP 최종 편집을 고려하여 A4, 한국어 논문형 글꼴, 넉넉한 줄간격을 적용한다.
"""

from __future__ import annotations

import html
import os
import re
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


WORKSPACE = Path(__file__).resolve().parents[1]
SRC = WORKSPACE / "논문_통합본_수정본.md"
OUT_DIR = WORKSPACE / "outputs" / "thesis_docx"
RUN_ID = datetime.now().strftime("%Y%m%d_%H%M%S")
OUT = OUT_DIR / f"논문_통합본_수정본_워드변환_{RUN_ID}.docx"

BODY_FONT = "함초롬바탕"
ALT_BODY_FONT = "Batang"
HEADING_FONT = "함초롬바탕"
CODE_FONT = "Consolas"


def set_run_font(run, font_name: str, size_pt: float | None = None, bold: bool | None = None):
    """한국어 글꼴이 깨지지 않도록 동아시아 글꼴까지 함께 지정한다."""
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), font_name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold


def set_paragraph(paragraph, *, align=None, before=0, after=4, line=1.75, first_line=False):
    if align is not None:
        paragraph.alignment = align
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Cm(0.7)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="808080", size="4"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_page_number(section):
    footer = section.footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, line=1.0)
    run = paragraph.add_run()
    set_run_font(run, BODY_FONT, 9)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_section(section):
    # 한글 논문 편집으로 넘기기 좋은 A4 기본값. 최종 여백은 HWP 조판에서 미세조정한다.
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(30)
    section.bottom_margin = Mm(25)
    section.left_margin = Mm(30)
    section.right_margin = Mm(30)
    section.header_distance = Mm(12)
    section.footer_distance = Mm(12)


def configure_styles(doc: Document):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.8)
    normal.paragraph_format.line_spacing = 1.75
    normal.paragraph_format.space_after = Pt(4)

    heading_specs = [
        ("Heading 1", 16, 18, 12, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 13, 14, 7, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 11.5, 10, 5, WD_ALIGN_PARAGRAPH.LEFT),
    ]
    for style_name, size, before, after, align in heading_specs:
        style = styles[style_name]
        style.font.name = HEADING_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.45
        style.paragraph_format.alignment = align

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(10.8)
        style.paragraph_format.line_spacing = 1.55
        style.paragraph_format.space_after = Pt(2)


def clean_text(text: str) -> str:
    text = html.unescape(text)
    text = text.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    text = re.sub(r"<sup>(.*?)</sup>", r"\1", text)
    text = re.sub(r"<sub>(.*?)</sub>", r"\1", text)
    text = re.sub(r"<[^>]+>", "", text)
    text = text.replace("&nbsp;", " ")
    return text.strip()


def add_inline_runs(paragraph, text: str, *, font=BODY_FONT, size=10.8, default_bold=False):
    """굵게, 인라인 코드 정도만 보존하는 가벼운 Markdown inline 처리."""
    text = clean_text(text)
    if not text:
        return
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        bold = default_bold
        run_font = font
        if part.startswith("**") and part.endswith("**"):
            part = part[2:-2]
            bold = True
        elif part.startswith("`") and part.endswith("`"):
            part = part[1:-1]
            run_font = CODE_FONT
        for idx, sub in enumerate(part.split("\n")):
            if idx:
                paragraph.add_run().add_break()
            run = paragraph.add_run(sub)
            set_run_font(run, run_font, size, bold)


def add_body_paragraph(doc, text: str, *, align=None, first_line=True):
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, align=align, before=0, after=4, line=1.75, first_line=first_line)
    add_inline_runs(paragraph, text)
    return paragraph


def add_heading(doc, level: int, text: str):
    style = "Heading 1" if level == 1 else "Heading 2" if level == 2 else "Heading 3"
    paragraph = doc.add_paragraph(style=style)
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline_runs(paragraph, text, font=HEADING_FONT, size=16 if level == 1 else 13 if level == 2 else 11.5, default_bold=True)
    return paragraph


def is_table_separator(line: str) -> bool:
    cells = [c.strip() for c in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip().startswith("|") or "|" not in line.strip()[1:]:
            break
        if not is_table_separator(line):
            rows.append([clean_text(c.strip()) for c in line.strip().strip("|").split("|")])
        i += 1
    return rows, i


def add_markdown_table(doc, rows: list[list[str]]):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    set_table_borders(table)

    for r_idx, row in enumerate(rows):
        for c_idx in range(ncols):
            cell = table.cell(r_idx, c_idx)
            value = row[c_idx] if c_idx < len(row) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, "E8EEF5")
            p = cell.paragraphs[0]
            set_paragraph(p, before=0, after=0, line=1.25, first_line=False)
            # 숫자·연도·짧은 항목은 가운데 정렬, 설명형 텍스트는 왼쪽 정렬한다.
            if len(value) <= 12 or re.fullmatch(r"[\d,.\-%~+() ]+", value):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline_runs(p, value, size=9.0 if ncols >= 6 else 9.6, default_bold=(r_idx == 0))

    # 긴 표가 페이지를 넘을 때 헤더행 반복을 허용한다.
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

    spacer = doc.add_paragraph()
    set_paragraph(spacer, before=0, after=4, line=1.0, first_line=False)


def resolve_image_path(raw_path: str) -> Path:
    path = raw_path.strip().strip("<>").replace("\\", "/")
    if path.startswith("file://"):
        path = path.replace("file://", "", 1)
    p = Path(path)
    if not p.is_absolute():
        p = WORKSPACE / p
    return p


def parse_image(line: str) -> tuple[str, Path] | None:
    # alt 텍스트에 괄호가 있어도 마지막 닫는 괄호 전까지 경로로 처리한다.
    m = re.match(r"!\[(.*?)\]\((.*)\)\s*$", line.strip())
    if not m:
        return None
    alt = clean_text(m.group(1))
    path = resolve_image_path(m.group(2))
    return alt, path


def add_image(doc, alt: str, path: Path, report: dict):
    if path.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2, line=1.0, first_line=False)
        run = paragraph.add_run()
        # A4 본문폭 안에서 그림을 안정적으로 배치한다.
        run.add_picture(str(path), width=Cm(14.6))
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph(caption, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=8, line=1.25, first_line=False)
        add_inline_runs(caption, alt, size=9.2, default_bold=False)
        report["inserted_images"].append(str(path.relative_to(WORKSPACE)))
    else:
        paragraph = doc.add_paragraph()
        set_paragraph(paragraph, before=4, after=4, line=1.4, first_line=False)
        add_inline_runs(paragraph, f"[그림 삽입 필요: {alt} / 경로: {path}]", size=10.2, default_bold=True)
        report["missing_images"].append(str(path))


def add_code_block(doc, code_lines: list[str]):
    if not code_lines:
        return
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph, before=4, after=6, line=1.2, first_line=False)
    run = paragraph.add_run("\n".join(code_lines).rstrip())
    set_run_font(run, CODE_FONT, 9.2, False)
    paragraph.paragraph_format.left_indent = Cm(0.5)


def should_skip_line(line: str, in_comment: bool) -> tuple[bool, bool]:
    stripped = line.strip()
    if stripped.startswith("<!--"):
        if "-->" not in stripped:
            return True, True
        return True, False
    if in_comment:
        return True, "-->" not in stripped
    if stripped in {"<div align=\"center\">", "</div>"}:
        return True, in_comment
    return False, in_comment


def build_docx():
    if not SRC.exists():
        raise FileNotFoundError(f"원본 Markdown을 찾을 수 없습니다: {SRC}")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    configure_section(doc.sections[0])
    add_page_number(doc.sections[0])
    configure_styles(doc)

    report = {"inserted_images": [], "missing_images": [], "tables": 0, "paragraphs": 0}

    in_comment = False
    in_code = False
    code_lines: list[str] = []
    center_mode = False
    i = 0

    while i < len(lines):
        raw = lines[i].lstrip("\ufeff")
        stripped = raw.strip()

        skip, in_comment = should_skip_line(raw, in_comment)
        if skip:
            if stripped == "<div align=\"center\">":
                center_mode = True
            elif stripped == "</div>":
                center_mode = False
            i += 1
            continue

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(raw)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            # 표지·내표지 등 Markdown 구분선은 Word 쪽나눔으로 변환한다.
            doc.add_page_break()
            i += 1
            continue

        image = parse_image(stripped)
        if image:
            add_image(doc, image[0], image[1], report)
            i += 1
            continue

        if stripped.startswith("|") and "|" in stripped[1:]:
            rows, next_i = parse_table(lines, i)
            if rows:
                add_markdown_table(doc, rows)
                report["tables"] += 1
                i = next_i
                continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            add_heading(doc, level, heading.group(2))
            i += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            set_paragraph(p, before=0, after=2, line=1.55, first_line=False)
            add_inline_runs(p, numbered.group(1))
        elif bullet:
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph(p, before=0, after=2, line=1.55, first_line=False)
            add_inline_runs(p, bullet.group(1))
        else:
            align = WD_ALIGN_PARAGRAPH.CENTER if center_mode else None
            first_line = not center_mode and not stripped.startswith("[") and not stripped.startswith("※")
            add_body_paragraph(doc, stripped, align=align, first_line=first_line)
            report["paragraphs"] += 1

        i += 1

    if code_lines:
        add_code_block(doc, code_lines)

    doc.core_properties.title = "전력설비 투자계획 수립을 위한 다기준 의사결정 기반 포트폴리오 최적화 프레임워크 연구"
    doc.core_properties.subject = "전력설비 투자계획 박사논문 Word 변환본"
    doc.core_properties.author = "Codex"
    doc.save(OUT)

    print(f"OUT={OUT}")
    print(f"INSERTED_IMAGES={len(report['inserted_images'])}")
    print(f"MISSING_IMAGES={len(report['missing_images'])}")
    print(f"TABLES={report['tables']}")
    print(f"PARAGRAPHS={report['paragraphs']}")
    if report["missing_images"]:
        print("MISSING_IMAGE_LIST=" + "|".join(report["missing_images"]))


if __name__ == "__main__":
    try:
        build_docx()
    except Exception as exc:
        print(f"오류: {exc}", file=sys.stderr)
        raise
