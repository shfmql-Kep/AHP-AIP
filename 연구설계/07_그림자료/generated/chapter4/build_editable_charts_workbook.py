from pathlib import Path
from openpyxl import load_workbook
import xlsxwriter


ROOT = Path(__file__).resolve().parents[3]
RESULT_FILE = (
    ROOT
    / "06_시뮬레이션결과"
    / "본문_시뮬레이션"
    / "simulation_chapter_v2_matlab_20260624_094857"
    / "simulation_chapter_v2_results.xlsx"
)
OUT_DIR = ROOT / "07_그림자료" / "excel"
OUT_DIR.mkdir(parents=True, exist_ok=True)
OUT_FILE = OUT_DIR / "그림4_2_및_연도별_SAIDI_편집용.xlsx"


def read_sheet_rows(workbook, sheet_name):
    ws = workbook[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    return [dict(zip(header, row)) for row in rows[1:]]


def method_rows(rows, method_name):
    return [row for row in rows if row.get("method") == method_name and row.get("asset_type_scope") == "all"]


def yearly_saidi_series(rows, method_name):
    data = {}
    for row in method_rows(rows, method_name):
        year = int(row["year"])
        data[year] = float(row["saidi_after_cumulative_min"])
    return data


def baseline_series(rows):
    data = {}
    for row in rows:
        if row.get("method") == "risk_greedy" and row.get("asset_type_scope") == "all":
            data[int(row["year"])] = float(row["baseline_saidi_min"])
    return data


def write_table(ws, start_row, start_col, headers, data, header_fmt, num_fmt=None):
    for c, header in enumerate(headers):
        ws.write(start_row, start_col + c, header, header_fmt)
    for r, row in enumerate(data, start_row + 1):
        for c, value in enumerate(row):
            fmt = num_fmt if isinstance(value, (int, float)) else None
            ws.write(r, start_col + c, value, fmt)


wb_src = load_workbook(RESULT_FILE, read_only=True, data_only=True)
value_annual = read_sheet_rows(wb_src, "03_value_combined_annual")
pi_annual = read_sheet_rows(wb_src, "05_pi_combined_annual")
integrated_annual = read_sheet_rows(wb_src, "06_integrated_annual")
wb_src.close()


# 그림 4.2 데이터: 단위는 억 원, 투자효율은 무차원
fig42_rows = [
    ["Risk Greedy", 1642.58, 759.98, 882.60, 2.161],
    ["IV Greedy", 1753.59, 760.70, 992.89, 2.305],
    ["IV ILP", 1786.34, 761.10, 1025.24, 2.347],
    ["IV GA", 1781.42, 761.66, 1019.76, 2.339],
]


years = [2026, 2027, 2028, 2029, 2030]
series_map = {
    "교체 없음": baseline_series(value_annual),
    "Risk Greedy": yearly_saidi_series(value_annual, "risk_greedy"),
    "IV Greedy": yearly_saidi_series(value_annual, "investment_value_greedy"),
    "IV ILP": yearly_saidi_series(value_annual, "investment_value_ilp"),
    "IV GA": yearly_saidi_series(value_annual, "investment_value_ga"),
    "PI Greedy": yearly_saidi_series(pi_annual, "pi_greedy"),
    "PI ILP": yearly_saidi_series(pi_annual, "pi_ilp"),
    "PI GA": yearly_saidi_series(pi_annual, "pi_ga"),
    "후보제한형 ILP": yearly_saidi_series(integrated_annual, "integrated_post_type_weight_ilp"),
    "통합형 ILP": yearly_saidi_series(integrated_annual, "integrated_pre_type_weight_ilp"),
}

saidi_headers = ["연도"] + list(series_map.keys())
saidi_rows = []
for year in years:
    saidi_rows.append([year] + [series_map[name][year] for name in series_map])


workbook = xlsxwriter.Workbook(str(OUT_FILE))
workbook.set_properties(
    {
        "title": "Figure 4.2 and Annual SAIDI Charts",
        "subject": "Editable thesis chart workbook",
        "author": "Codex",
        "comments": "Chart1.crtx 색상 체계를 참고하여 편집 가능한 엑셀 차트로 구성함",
    }
)

fmt_title = workbook.add_format({"bold": True, "font_name": "맑은 고딕", "font_size": 14, "font_color": "#1F1F1F"})
fmt_header = workbook.add_format(
    {
        "bold": True,
        "font_name": "맑은 고딕",
        "font_size": 10,
        "bg_color": "#D9EAF7",
        "border": 1,
        "align": "center",
        "valign": "vcenter",
    }
)
fmt_note = workbook.add_format({"font_name": "맑은 고딕", "font_size": 9, "font_color": "#666666"})
fmt_text = workbook.add_format({"font_name": "맑은 고딕", "font_size": 10, "border": 1})
fmt_num1 = workbook.add_format({"font_name": "맑은 고딕", "font_size": 10, "border": 1, "num_format": "#,##0.0"})
fmt_num3 = workbook.add_format({"font_name": "맑은 고딕", "font_size": 10, "border": 1, "num_format": "0.000"})
fmt_year = workbook.add_format({"font_name": "맑은 고딕", "font_size": 10, "border": 1, "align": "center"})


ws1 = workbook.add_worksheet("Fig4.2_Data")
ws1.hide_gridlines(2)
ws1.write("A1", "그림 4.2 데이터: 리스크 및 투자가치 기반 방식 전체 성과 비교", fmt_title)
ws1.write("A2", "단위: 금액 = 억 원, 투자효율 = 무차원", fmt_note)
write_table(
    ws1,
    3,
    0,
    ["기법", "리스크 저감량", "투자비용", "투자가치", "투자효율"],
    fig42_rows,
    fmt_header,
    fmt_num1,
)
for r in range(4, 8):
    ws1.write(r, 0, fig42_rows[r - 4][0], fmt_text)
    ws1.write_number(r, 4, fig42_rows[r - 4][4], fmt_num3)
ws1.set_column("A:A", 18)
ws1.set_column("B:D", 16)
ws1.set_column("E:E", 12)


chart1 = workbook.add_chart({"type": "column"})
chart1.add_series(
    {
        "name": "=Fig4.2_Data!$B$4",
        "categories": "=Fig4.2_Data!$A$5:$A$8",
        "values": "=Fig4.2_Data!$B$5:$B$8",
        "fill": {"color": "#FFC000"},
        "border": {"color": "#222222"},
        "data_labels": {"value": True, "num_format": "#,##0"},
    }
)
chart1.add_series(
    {
        "name": "=Fig4.2_Data!$C$4",
        "categories": "=Fig4.2_Data!$A$5:$A$8",
        "values": "=Fig4.2_Data!$C$5:$C$8",
        "fill": {"color": "#5B9BD5"},
        "border": {"color": "#222222"},
        "data_labels": {"value": True, "num_format": "#,##0"},
    }
)
chart1.add_series(
    {
        "name": "=Fig4.2_Data!$D$4",
        "categories": "=Fig4.2_Data!$A$5:$A$8",
        "values": "=Fig4.2_Data!$D$5:$D$8",
        "fill": {"color": "#ED7D31"},
        "border": {"color": "#222222"},
        "data_labels": {"value": True, "num_format": "#,##0"},
    }
)

chart1_line = workbook.add_chart({"type": "line"})
chart1_line.add_series(
    {
        "name": "=Fig4.2_Data!$E$4",
        "categories": "=Fig4.2_Data!$A$5:$A$8",
        "values": "=Fig4.2_Data!$E$5:$E$8",
        "y2_axis": True,
        "line": {"color": "#0000FF", "width": 2.5, "dash_type": "dash"},
        "marker": {"type": "circle", "size": 7, "border": {"color": "#0000FF"}, "fill": {"color": "#0000FF"}},
        "data_labels": {"value": True, "num_format": "0.00", "font": {"color": "#0000FF", "bold": True}},
    }
)
chart1.combine(chart1_line)
chart1.set_title({"name": ""})
chart1.set_legend({"position": "top"})
chart1.set_x_axis({"name": "최적화 기법", "name_font": {"bold": True}, "num_font": {"rotation": 0}})
chart1.set_y_axis({"name": "금액 [억 원]", "major_gridlines": {"visible": True, "line": {"color": "#D9D9D9"}}, "min": 0, "max": 2000, "num_format": "#,##0"})
chart1.set_y2_axis({"name": "투자효율 [-]", "min": 0, "max": 3.0, "num_format": "0.0", "name_font": {"color": "#0000FF"}, "num_font": {"color": "#0000FF"}})
chart1.set_chartarea({"border": {"color": "#222222", "width": 1}})
chart1.set_plotarea({"border": {"color": "#222222", "width": 1}, "fill": {"color": "#FFFFFF"}})
chart1.set_size({"width": 920, "height": 540})
ws1.insert_chart("G3", chart1)


ws2 = workbook.add_worksheet("SAIDI_Data")
ws2.hide_gridlines(2)
ws2.write("A1", "연도별 SAIDI 변화 데이터", fmt_title)
ws2.write("A2", "값은 누적 교체효과를 반영한 투자 후 잔여 SAIDI이며, '교체 없음'은 기준 SAIDI이다.", fmt_note)
write_table(ws2, 3, 0, saidi_headers, saidi_rows, fmt_header, fmt_num3)
for r, row in enumerate(saidi_rows, 4):
    ws2.write_number(r, 0, row[0], fmt_year)
    for c, value in enumerate(row[1:], 1):
        ws2.write_number(r, c, value, fmt_num3)
ws2.set_column("A:A", 10)
ws2.set_column("B:K", 15)


chart2 = workbook.add_chart({"type": "line"})
line_colors = {
    "교체 없음": "#A6A6A6",
    "Risk Greedy": "#FFC000",
    "IV Greedy": "#70AD47",
    "IV ILP": "#ED7D31",
    "IV GA": "#F4B183",
    "PI Greedy": "#9DC3E6",
    "PI ILP": "#5B9BD5",
    "PI GA": "#2F75B5",
    "후보제한형 ILP": "#7030A0",
    "통합형 ILP": "#0000FF",
}
for idx, name in enumerate(series_map.keys(), start=1):
    color = line_colors.get(name, "#333333")
    dash_type = "dash" if name in {"교체 없음", "Risk Greedy"} else "solid"
    width = 2.25 if name in {"IV ILP", "PI ILP", "통합형 ILP", "후보제한형 ILP"} else 1.5
    chart2.add_series(
        {
            "name": f"=SAIDI_Data!${chr(65+idx)}$4",
            "categories": "=SAIDI_Data!$A$5:$A$9",
            "values": f"=SAIDI_Data!${chr(65+idx)}$5:${chr(65+idx)}$9",
            "line": {"color": color, "width": width, "dash_type": dash_type},
            "marker": {"type": "circle", "size": 5, "border": {"color": color}, "fill": {"color": color}},
        }
    )
chart2.set_title({"name": ""})
chart2.set_legend({"position": "bottom"})
chart2.set_x_axis({"name": "연도", "name_font": {"bold": True}})
chart2.set_y_axis({"name": "SAIDI [분/고객·년]", "major_gridlines": {"visible": True, "line": {"color": "#D9D9D9"}}, "min": 0.6, "max": 2.1, "num_format": "0.0"})
chart2.set_chartarea({"border": {"color": "#222222", "width": 1}})
chart2.set_plotarea({"border": {"color": "#222222", "width": 1}, "fill": {"color": "#FFFFFF"}})
chart2.set_size({"width": 920, "height": 540})
ws2.insert_chart("M3", chart2)


ws3 = workbook.add_worksheet("SAIDI_Key_Methods")
ws3.hide_gridlines(2)
ws3.write("A1", "연도별 SAIDI 변화: 주요 기법 비교", fmt_title)
ws3.write("A2", "논문 본문 그림으로 사용할 경우 선이 과밀하지 않도록 주요 기법만 표시한다.", fmt_note)
key_methods = ["교체 없음", "Risk Greedy", "IV ILP", "PI ILP", "후보제한형 ILP", "통합형 ILP"]
key_headers = ["연도"] + key_methods
key_rows = [[year] + [series_map[name][year] for name in key_methods] for year in years]
write_table(ws3, 3, 0, key_headers, key_rows, fmt_header, fmt_num3)
for r, row in enumerate(key_rows, 4):
    ws3.write_number(r, 0, row[0], fmt_year)
    for c, value in enumerate(row[1:], 1):
        ws3.write_number(r, c, value, fmt_num3)
ws3.set_column("A:A", 10)
ws3.set_column("B:G", 16)

chart3 = workbook.add_chart({"type": "line"})
for idx, name in enumerate(key_methods, start=1):
    color = line_colors.get(name, "#333333")
    dash_type = "dash" if name in {"교체 없음", "Risk Greedy"} else "solid"
    width = 2.75 if name in {"통합형 ILP", "후보제한형 ILP"} else 2.0
    chart3.add_series(
        {
            "name": f"=SAIDI_Key_Methods!${chr(65+idx)}$4",
            "categories": "=SAIDI_Key_Methods!$A$5:$A$9",
            "values": f"=SAIDI_Key_Methods!${chr(65+idx)}$5:${chr(65+idx)}$9",
            "line": {"color": color, "width": width, "dash_type": dash_type},
            "marker": {"type": "circle", "size": 6, "border": {"color": color}, "fill": {"color": color}},
        }
    )
chart3.set_title({"name": ""})
chart3.set_legend({"position": "bottom"})
chart3.set_x_axis({"name": "연도", "name_font": {"bold": True}})
chart3.set_y_axis({"name": "SAIDI [분/고객·년]", "major_gridlines": {"visible": True, "line": {"color": "#D9D9D9"}}, "min": 0.6, "max": 2.1, "num_format": "0.0"})
chart3.set_chartarea({"border": {"color": "#222222", "width": 1}})
chart3.set_plotarea({"border": {"color": "#222222", "width": 1}, "fill": {"color": "#FFFFFF"}})
chart3.set_size({"width": 920, "height": 540})
ws3.insert_chart("I3", chart3)


ws4 = workbook.add_worksheet("README")
ws4.hide_gridlines(2)
ws4.write("A1", "편집 안내", fmt_title)
readme = [
    ["항목", "내용"],
    ["Fig4.2_Data", "그림 4.2 편집용 데이터와 조합형 차트가 포함되어 있다."],
    ["SAIDI_Data", "전체 방법의 연도별 SAIDI 변화 데이터와 전체 방법 라인 차트가 포함되어 있다."],
    ["SAIDI_Key_Methods", "본문 삽입용으로 쓰기 쉬운 주요 기법 라인 차트가 포함되어 있다."],
    ["색상", "Chart1.crtx의 기본 계열색(노랑, 파랑, 주황, 파란 점선)을 반영하였다."],
    ["원천 파일", str(RESULT_FILE)],
]
write_table(ws4, 3, 0, readme[0], readme[1:], fmt_header, None)
for r in range(4, 4 + len(readme) - 1):
    ws4.write(r, 0, readme[r - 3][0], fmt_text)
    ws4.write(r, 1, readme[r - 3][1], fmt_text)
ws4.set_column("A:A", 22)
ws4.set_column("B:B", 120)

workbook.close()
print(OUT_FILE)
