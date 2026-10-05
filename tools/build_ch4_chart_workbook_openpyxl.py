from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
INPUT_JSON = ROOT / "tmp" / "ch4_simulation_chart_data.json"
OUTPUT_DIR = ROOT / "연구설계" / "06_시뮬레이션결과" / "그림작성용"
OUTPUT_XLSX = OUTPUT_DIR / "최적화시뮬레이션_그림작성용_데이터.xlsx"
INSPECT_NDJSON = OUTPUT_DIR / "최적화시뮬레이션_그림작성용_데이터.xlsx.inspect.ndjson"


PREFERRED_HEADERS = {
    "Chart_Guide": ["item", "value", "note"],
    "README": ["item", "value"],
    "Final_Comparison": [
        "method",
        "method_label",
        "selected_count",
        "investment_cost_eok",
        "risk_reduction_eok",
        "investment_value_eok",
        "saidi_reduction_min",
        "asset_level_pi",
        "integrated_pi",
        "investment_efficiency",
    ],
    "All_Method_Total": [
        "method",
        "method_label",
        "selected_count",
        "investment_cost_eok",
        "risk_reduction_eok",
        "investment_value_eok",
        "saidi_reduction_min",
        "asset_level_pi",
        "integrated_pi",
        "investment_efficiency",
    ],
    "Baseline_vs_RiskGreedy": [
        "method_label",
        "selected_count",
        "count_vs_baseline",
        "investment_cost_eok",
        "cost_vs_baseline",
        "risk_reduction_eok",
        "risk_reduction_vs_baseline",
        "investment_value_eok",
        "investment_value_vs_baseline",
        "saidi_reduction_min",
        "saidi_vs_baseline",
        "asset_level_pi",
        "asset_level_pi_vs_baseline",
        "integrated_pi",
        "integrated_pi_vs_baseline",
        "investment_efficiency",
    ],
    "AssetLevelPI_Type_Annual": [
        "asset_type_label",
        "method_label",
        "year",
        "selected_count",
        "investment_cost_eok",
        "risk_reduction_eok",
        "investment_value_eok",
        "saidi_reduction_min",
        "asset_level_pi",
        "integrated_pi",
        "investment_efficiency",
    ],
    "Missing_Check": [
        "asset_type_label",
        "asset_level_pi_total_rows",
        "asset_level_pi_annual_rows",
        "asset_level_pi_selected_asset_rows",
        "check_result",
    ],
}


SHEET_ORDER = [
    "Chart_Guide",
    "README",
    "Final_Comparison",
    "All_Method_Total",
    "Baseline_vs_RiskGreedy",
    "Chart_Combo_대표",
    "Chart_Combo_전체",
    "Yearly_Cost",
    "Yearly_Count",
    "Yearly_IV",
    "Value_Type_Total",
    "AssetLevelPI_Type_Total",
    "AssetLevelPI_IV_vs_PI",
    "AssetLevelPI_Type_Annual",
    "Integrated_Type_Total",
    "Missing_Check",
]


NUMBER_FORMATS = {
    "selected_count": "#,##0",
    "year": "0",
    "investment_cost_eok": "#,##0.00",
    "risk_reduction_eok": "#,##0.00",
    "investment_value_eok": "#,##0.00",
    "saidi_reduction_min": "0.0000",
    "asset_level_pi": "#,##0.00",
    "integrated_pi": "#,##0.00",
    "investment_efficiency": "0.00",
    "count_vs_baseline": "0.0%",
    "cost_vs_baseline": "0.0%",
    "risk_reduction_vs_baseline": "0.0%",
    "investment_value_vs_baseline": "0.0%",
    "saidi_vs_baseline": "0.0%",
    "asset_level_pi_vs_baseline": "0.0%",
    "integrated_pi_vs_baseline": "0.0%",
}


METHOD_ORDER = [
    "Risk 그리디",
    "투자가치 그리디",
    "투자가치 ILP",
    "투자가치 GA",
    "PI 그리디",
    "PI ILP",
    "PI GA",
    "후보제한형 ILP",
    "후보제한형 GA",
    "통합형 ILP",
    "통합형 GA",
]


def union_headers(rows: list[dict], preferred: list[str] | None = None) -> list[str]:
    headers: list[str] = []
    seen: set[str] = set()
    for header in preferred or []:
        if header not in seen:
            headers.append(header)
            seen.add(header)
    for row in rows:
        for header in row:
            if header not in seen:
                headers.append(header)
                seen.add(header)
    return headers


def by_method_label(rows: list[dict]) -> dict[str, dict]:
    return {str(row.get("method_label")): row for row in rows if row.get("method_label")}


def pct(value, base):
    if value is None or base in (None, 0):
        return None
    return (float(value) - float(base)) / abs(float(base))


def baseline_comparison(rows: list[dict], baseline_label: str = "Risk 그리디") -> list[dict]:
    row_map = by_method_label(rows)
    base = row_map.get(baseline_label)
    if not base:
        return []

    def sort_key(row: dict) -> int:
        label = str(row.get("method_label", ""))
        return METHOD_ORDER.index(label) if label in METHOD_ORDER else 999

    out = []
    for row in sorted(rows, key=sort_key):
        out.append(
            {
                "method_label": row.get("method_label"),
                "selected_count": row.get("selected_count"),
                "count_vs_baseline": pct(row.get("selected_count"), base.get("selected_count")),
                "investment_cost_eok": row.get("investment_cost_eok"),
                "cost_vs_baseline": pct(row.get("investment_cost_eok"), base.get("investment_cost_eok")),
                "risk_reduction_eok": row.get("risk_reduction_eok"),
                "risk_reduction_vs_baseline": pct(row.get("risk_reduction_eok"), base.get("risk_reduction_eok")),
                "investment_value_eok": row.get("investment_value_eok"),
                "investment_value_vs_baseline": pct(row.get("investment_value_eok"), base.get("investment_value_eok")),
                "saidi_reduction_min": row.get("saidi_reduction_min"),
                "saidi_vs_baseline": pct(row.get("saidi_reduction_min"), base.get("saidi_reduction_min")),
                "asset_level_pi": row.get("asset_level_pi"),
                "asset_level_pi_vs_baseline": pct(row.get("asset_level_pi"), base.get("asset_level_pi")),
                "integrated_pi": row.get("integrated_pi"),
                "integrated_pi_vs_baseline": pct(row.get("integrated_pi"), base.get("integrated_pi")),
                "investment_efficiency": row.get("investment_efficiency"),
            }
        )
    return out


def aggregate_totals(rows: list[dict]) -> list[dict]:
    groups: dict[str, dict] = {}
    for row in rows:
        key = str(row.get("method") or row.get("method_label"))
        if not key:
            continue
        if key not in groups:
            groups[key] = {
                "method": row.get("method", key),
                "method_label": row.get("method_label", key),
                "selected_count": 0,
                "investment_cost_eok": 0.0,
                "risk_reduction_eok": 0.0,
                "investment_value_eok": 0.0,
                "saidi_reduction_min": 0.0,
                "asset_level_pi": 0.0,
                "integrated_pi": 0.0,
            }
        item = groups[key]
        for field in [
            "selected_count",
            "investment_cost_eok",
            "risk_reduction_eok",
            "investment_value_eok",
            "saidi_reduction_min",
            "asset_level_pi",
            "integrated_pi",
        ]:
            item[field] += float(row.get(field) or 0)
    out = []
    for row in groups.values():
        cost = row["investment_cost_eok"]
        row["selected_count"] = round(row["selected_count"])
        row["investment_efficiency"] = None if cost == 0 else row["risk_reduction_eok"] / cost
        out.append(row)

    def sort_key(row: dict) -> int:
        label = str(row.get("method_label", ""))
        return METHOD_ORDER.index(label) if label in METHOD_ORDER else 999

    return sorted(out, key=sort_key)


def filter_by_method_label(rows: list[dict], labels: list[str]) -> list[dict]:
    label_set = set(labels)
    return [row for row in rows if row.get("method_label") in label_set]


def guide_rows(source_file: str) -> list[dict]:
    return [
        {
            "item": "source_workbook",
            "value": source_file,
            "note": "최신 본문 시뮬레이션 결과 파일에서 그림 작성용 데이터를 재추출하였다.",
        },
        {
            "item": "representative_methods",
            "value": "Risk 그리디, 투자가치 ILP, PI ILP, 통합형 ILP",
            "note": "통합형 ILP를 최종 대표 통합 방식으로 사용한다.",
        },
        {
            "item": "Chart_Combo",
            "value": "Chart_Combo_대표 또는 Chart_Combo_전체 사용",
            "note": "막대: Risk 저감량·투자비용·투자가치, 보조축: 투자효율",
        },
        {
            "item": "Yearly_Cost",
            "value": "연도별 투자비용 비교",
            "note": "통합형 ILP 기준으로 갱신하였다.",
        },
        {
            "item": "Yearly_Count",
            "value": "연도별 투자대수 비교",
            "note": "통합형 ILP 기준으로 갱신하였다.",
        },
        {
            "item": "AssetLevelPI_IV_vs_PI",
            "value": "설비유형별 필터 사용",
            "note": "투자가치 ILP와 PI 기반 방법을 설비유형별로 비교한다.",
        },
        {
            "item": "Missing_Check",
            "value": "설비 단위 PI 출력 누락 점검",
            "note": "지중개폐기와 가공배전선로가 최신 설비 단위 PI 결과에 포함된다.",
        },
    ]


def write_sheet(wb: Workbook, name: str, rows: list[dict], preferred: list[str] | None = None) -> None:
    ws = wb.create_sheet(title=name[:31])
    headers = union_headers(rows, preferred)
    if not headers:
        headers = ["empty"]
        rows = [{"empty": ""}]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(header) for header in headers])

    header_fill = PatternFill("solid", fgColor="0B2A5B")
    header_font = Font(name="맑은 고딕", color="FFFFFF", bold=True)
    body_font = Font(name="맑은 고딕", size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            number_format = NUMBER_FORMATS.get(headers[cell.column - 1])
            if number_format:
                cell.number_format = number_format

    for col_idx, header in enumerate(headers, start=1):
        values = [str(header)]
        for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, max_row=ws.max_row):
            values.extend("" if item.value is None else str(item.value) for item in cell)
        width = min(max(len(v) for v in values) + 2, 42)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

    table_ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    table_name = "".join(ch if ch.isalnum() else "_" for ch in name)[:220] + "_Table"
    table = Table(displayName=table_name, ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)


def write_inspect(wb: Workbook, output_path: Path) -> None:
    lines = []
    for ws in wb.worksheets:
        values = []
        max_preview_row = min(ws.max_row, 6)
        max_preview_col = min(ws.max_column, 8)
        for row in ws.iter_rows(min_row=1, max_row=max_preview_row, max_col=max_preview_col, values_only=True):
            values.append(list(row))
        lines.append(
            json.dumps(
                {
                    "kind": "sheet",
                    "name": ws.title,
                    "rows": ws.max_row,
                    "cols": ws.max_column,
                    "preview": values,
                },
                ensure_ascii=False,
            )
        )
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    data = json.loads(INPUT_JSON.read_text(encoding="utf-8"))
    sheets = data["sheets"]

    all_method_totals = aggregate_totals(
        [
            *sheets.get("Value_Type_Total", []),
            *sheets.get("LocalPI_Type_Total", []),
            *sheets.get("Integrated_Type_Total", []),
        ]
    )
    representative_totals = filter_by_method_label(
        all_method_totals,
        ["Risk 그리디", "투자가치 ILP", "PI ILP", "통합형 ILP"],
    )

    prepared = {
        **sheets,
        "Chart_Guide": guide_rows(data["source_file"]),
        "All_Method_Total": all_method_totals,
        "Baseline_vs_RiskGreedy": baseline_comparison(all_method_totals),
        "Chart_Combo_대표": representative_totals,
        "Chart_Combo_전체": all_method_totals,
        "AssetLevelPI_Type_Total": sheets.get("LocalPI_Type_Total", []),
        "AssetLevelPI_IV_vs_PI": sheets.get("LocalPI_IV_vs_PI", []),
        "AssetLevelPI_Type_Annual": sheets.get("LocalPI_Type_Annual", []),
    }

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if OUTPUT_XLSX.exists():
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        backup_path = OUTPUT_XLSX.with_name(f"{OUTPUT_XLSX.stem}_backup_{stamp}{OUTPUT_XLSX.suffix}")
        shutil.copy2(OUTPUT_XLSX, backup_path)
        print(f"backup={backup_path}")

    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name in SHEET_ORDER:
        rows = prepared.get(sheet_name, [])
        write_sheet(wb, sheet_name, rows, PREFERRED_HEADERS.get(sheet_name))

    wb.save(OUTPUT_XLSX)
    write_inspect(wb, INSPECT_NDJSON)
    print(f"output={OUTPUT_XLSX}")
    print(f"inspect={INSPECT_NDJSON}")


if __name__ == "__main__":
    main()
