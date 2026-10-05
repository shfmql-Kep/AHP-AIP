from pathlib import Path
import re
import html

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "연구설계" / "01_논문_장별" / "10_부록_설문지.md"
OUT_DIR = ROOT / "연구설계" / "03_설문조사" / "워드붙여넣기용"
OUT_DOCX = OUT_DIR / "부록_설문지_평가기준_및_문항_워드붙여넣기용.docx"

DOC_FONT = "바탕"
HEAD_FONT = "바탕"


def set_run_font(run, name=DOC_FONT, size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def clean_inline(text):
    text = html.unescape(text)
    text = text.replace("&nbsp;", " ")
    text = re.sub(r"<br\s*/?>", "\n", text)
    text = text.replace("**", "")
    text = text.replace("*", "")
    text = text.replace("`", "")
    text = text.replace("←←", "←")
    text = text.replace("→→", "→")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def add_rich_paragraph(doc, line, style=None):
    line = html.unescape(line).replace("&nbsp;", " ")
    if line.startswith("> "):
        line = line[2:]
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.15
    parts = re.split(r"(\*\*.*?\*\*)", line)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        content = part[2:-2] if bold else part
        content = content.replace("*", "")
        run = paragraph.add_run(content)
        set_run_font(run, size=10.5, bold=bold)
    return paragraph


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table, rows):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    col_count = len(rows[0]) if rows else 0
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "D9D9D9")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=9.2 if col_count >= 5 else 10, bold=(row_index == 0))


def add_markdown_table(doc, table_lines):
    rows = []
    for line in table_lines:
        if re.match(r"^\|\s*:?-{3,}:?", line):
            continue
        if not line.strip().startswith("|"):
            continue
        cells = [clean_inline(cell) for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    if not rows:
        return
    max_cols = max(len(row) for row in rows)
    rows = [row + [""] * (max_cols - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=max_cols)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value
    style_table(table, rows)
    doc.add_paragraph()


def build_docx():
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    source_text = SOURCE.read_text(encoding="utf-8")
    start = source_text.index("## II. 평가기준 체계 안내")
    end = source_text.index("## V. 설문 결과 요약")
    lines = source_text[start:end].strip().splitlines()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = DOC_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), DOC_FONT)
    normal.font.size = Pt(10.5)
    for style_name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11)]:
        style = styles[style_name]
        style.font.name = HEAD_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("부록 설문지: 평가기준 체계 및 설문 문항")
    set_run_font(title_run, size=14, bold=True)
    doc.add_paragraph()

    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip() or line.strip() == "---":
            index += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(doc, table_lines)
            continue
        if line.startswith("## "):
            paragraph = doc.add_paragraph(clean_inline(line[3:]), style="Heading 1")
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(6)
        elif line.startswith("### "):
            paragraph = doc.add_paragraph(clean_inline(line[4:]), style="Heading 2")
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)
        elif line.startswith("> - "):
            paragraph = doc.add_paragraph(style="List Bullet")
            run = paragraph.add_run(clean_inline(line[4:]))
            set_run_font(run, size=10)
        elif line.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.4)
            paragraph.paragraph_format.space_after = Pt(4)
            run = paragraph.add_run(clean_inline(line[2:]))
            set_run_font(run, size=10)
        elif line.startswith("**") and line.endswith("**"):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(clean_inline(line))
            set_run_font(run, size=10.5, bold=True)
        else:
            add_rich_paragraph(doc, line)
        index += 1

    doc.core_properties.title = "부록 설문지 평가기준 및 문항"
    doc.core_properties.author = ""
    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    print(build_docx())
